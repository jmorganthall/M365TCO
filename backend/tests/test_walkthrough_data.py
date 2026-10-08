"""The walkthrough's data (docs/WALKTHROUGH.md §12, TARGET_SCHEMA D19–D21): Microsoft
renewal dates, the unused-seat answer and coverage-gap answers — stored, copied,
cleaned up, and read by the timed headline and the readouts."""

import io
from datetime import date, timedelta

from openpyxl import load_workbook


def _eng(client, name="Walkthrough Co"):
    eng = client.post("/api/engagements", json={"customer_name": name}).json()
    return eng["id"], date.fromisoformat(eng["workshop_date"])


def _persona(client, eid, name="KW", headcount=100):
    return client.post(f"/api/engagements/{eid}/personas",
                       json={"name": name, "headcount": headcount}).json()["id"]


def _outcome(client, eid, seed_key):
    return next(o for o in client.get(f"/api/engagements/{eid}/outcomes").json()
                if o["seed_key"] == seed_key)["id"]


def _headline(client, eid):
    return client.post(f"/api/engagements/{eid}/compute").json()["rollup"]["headline"]


# ---- Coverage gap answers -------------------------------------------------

def test_gap_answer_is_one_per_persona_and_outcome(client):
    eid, _ = _eng(client)
    pid, oid = _persona(client, eid), _outcome(client, eid, "endpoint-edr")
    url = f"/api/engagements/{eid}/coverage-gap-answers"
    first = client.put(url, json={"persona_id": pid, "outcome_id": oid,
                                  "answer": "NotDeliveredToday"}).json()
    second = client.put(url, json={"persona_id": pid, "outcome_id": oid,
                                   "answer": "CoveredOutsideInventory"}).json()
    assert second["id"] == first["id"]                       # replaced, not duplicated
    rows = client.get(url).json()
    assert [(r["persona_id"], r["outcome_id"], r["answer"]) for r in rows] == \
        [(pid, oid, "CoveredOutsideInventory")]
    assert client.delete(f"{url}/{first['id']}").status_code == 204
    assert client.get(url).json() == []


def test_gap_answer_rejects_another_engagements_rows_and_bad_answers(client):
    eid, _ = _eng(client, "Mine Co")
    other, _ = _eng(client, "Other Co")
    pid, oid = _persona(client, other), _outcome(client, other, "endpoint-edr")
    url = f"/api/engagements/{eid}/coverage-gap-answers"
    assert client.put(url, json={"persona_id": pid, "outcome_id": oid,
                                 "answer": "NotDeliveredToday"}).status_code == 404
    mine = _persona(client, eid)
    assert client.put(url, json={"persona_id": mine,
                                 "outcome_id": _outcome(client, eid, "endpoint-edr"),
                                 "answer": "Maybe"}).status_code == 422


def _gap_engagement(client):
    """KW holds Microsoft 365 E3 and moves to E5, which adds capabilities nothing
    delivers today."""
    eid, _ = _eng(client, "Gap Co")
    pid = _persona(client, eid)
    client.post(f"/api/engagements/{eid}/current-licenses", json={
        "sku_reference": "Microsoft 365 E3", "quantity_assigned": 100,
        "unit_price_paid_annual": 400, "persona_ids": [pid]})
    client.post(f"/api/engagements/{eid}/scenarios", json={
        "persona_id": pid, "target_sku_reference": "Microsoft 365 E5",
        "target_unit_price_annual": 600, "in_scope": True})
    return eid, pid


def test_covered_outside_counts_as_delivered_today_and_never_as_new(client):
    eid, pid = _gap_engagement(client)
    gaps = client.get(f"/api/engagements/{eid}/coverage-gaps").json()["personas"][0]
    edr = next(o for o in gaps["uncovered_outcomes"] if o["id"] == _outcome(client, eid, "endpoint-edr"))
    assert edr["answer"] is None
    client.put(f"/api/engagements/{eid}/coverage-gap-answers", json={
        "persona_id": pid, "outcome_id": edr["id"], "answer": "CoveredOutsideInventory"})
    gaps = client.get(f"/api/engagements/{eid}/coverage-gaps").json()["personas"][0]
    assert edr["id"] not in [o["id"] for o in gaps["uncovered_outcomes"]]
    assert [o["id"] for o in gaps["covered_outside_outcomes"]] == [edr["id"]]
    result = client.post(f"/api/engagements/{eid}/compute").json()
    entry = result["new_outcomes"][0]
    assert edr["id"] not in [o["id"] for o in entry["outcomes"] + entry["unconfirmed_outcomes"]]


def test_deleting_a_persona_or_outcome_removes_its_answers(client):
    eid, pid = _gap_engagement(client)
    url = f"/api/engagements/{eid}/coverage-gap-answers"
    edr, epp = _outcome(client, eid, "endpoint-edr"), _outcome(client, eid, "endpoint-protection")
    for oid in (edr, epp):
        client.put(url, json={"persona_id": pid, "outcome_id": oid, "answer": "NotDeliveredToday"})
    client.delete(f"/api/engagements/{eid}/outcomes/{edr}")
    assert [r["outcome_id"] for r in client.get(url).json()] == [epp]
    client.delete(f"/api/engagements/{eid}/personas/{pid}")
    assert client.get(url).json() == []


def test_carve_out_inherits_the_parents_answers(client):
    eid, pid = _gap_engagement(client)
    edr = _outcome(client, eid, "endpoint-edr")
    client.put(f"/api/engagements/{eid}/coverage-gap-answers", json={
        "persona_id": pid, "outcome_id": edr, "answer": "CoveredOutsideInventory"})
    child = client.post(f"/api/engagements/{eid}/personas/{pid}/carve",
                        json={"seats": 20}).json()
    answers = {(r["persona_id"], r["outcome_id"]): r["answer"] for r in
               client.get(f"/api/engagements/{eid}/coverage-gap-answers").json()}
    assert answers[(child["id"], edr)] == "CoveredOutsideInventory"


def test_duplicate_keeps_dates_answers_and_entitlement_scope(client):
    eid, pid = _gap_engagement(client)
    client.patch(f"/api/engagements/{eid}", json={"microsoft_renewal_date": "2027-03-01"})
    lic = client.get(f"/api/engagements/{eid}/current-licenses").json()[0]
    client.patch(f"/api/engagements/{eid}/current-licenses/{lic['id']}", json={
        "quantity_purchased": 120, "renewal_date": "2027-06-01",
        "unused_seats_answer": "NotNeeded", "coverage_scope": "TenantWide"})
    client.put(f"/api/engagements/{eid}/coverage-gap-answers", json={
        "persona_id": pid, "outcome_id": _outcome(client, eid, "endpoint-edr"),
        "answer": "NotDeliveredToday"})
    src = client.get(f"/api/engagements/{eid}").json()
    copy = client.post(f"/api/engagements/{eid}/duplicate").json()
    assert copy["workshop_date"] == src["workshop_date"]
    assert copy["microsoft_renewal_date"] == "2027-03-01"
    clic = client.get(f"/api/engagements/{copy['id']}/current-licenses").json()[0]
    assert (clic["renewal_date"], clic["unused_seats_answer"], clic["coverage_scope"]) == \
        ("2027-06-01", "NotNeeded", "TenantWide")
    assert [r["answer"] for r in
            client.get(f"/api/engagements/{copy['id']}/coverage-gap-answers").json()] == \
        ["NotDeliveredToday"]
    # Same data, same numbers.
    assert _headline(client, copy["id"])["amount"] == _headline(client, eid)["amount"]


# ---- Renewal dates and unused seats through the API ------------------------

def test_overlicensing_counts_from_the_lines_renewal(client):
    eid, workshop = _eng(client, "Unused Co")
    client.patch(f"/api/engagements/{eid}",
                 json={"microsoft_renewal_date": (workshop + timedelta(days=1)).isoformat()})
    lines = {}
    for sku, bought, answer in (("E3", 120, "NotNeeded"), ("F3", 50, "Intended"), ("E5", 15, None)):
        lines[sku] = client.post(f"/api/engagements/{eid}/current-licenses", json={
            "sku_reference": sku, "quantity_purchased": bought,
            "quantity_assigned": bought - 20 if sku == "E3" else bought - 10,
            "unit_price_paid_annual": 240, "unused_seats_answer": answer}).json()
    h = _headline(client, eid)
    # Only E3's 20 not-needed seats count: $4,800/yr from month 1 → 35 months.
    assert h["overlicensing_amount"] == 14000.0
    assert [i["license_id"] for i in h["items"]] == [lines["E3"]["id"]]
    assert {u["sku_reference"]: u["answer"] for u in h["unused_seat_lines"]} == \
        {"E3": "NotNeeded", "F3": "Intended", "E5": None}
    # The line's own renewal overrides the agreement's.
    client.patch(f"/api/engagements/{eid}/current-licenses/{lines['E3']['id']}",
                 json={"renewal_date": workshop.isoformat()})
    assert _headline(client, eid)["overlicensing_amount"] == 14400.0      # all 36 months


def test_readouts_show_how_the_headline_is_timed(client):
    eid, _ = _eng(client, "Timed Readout Co")
    client.post(f"/api/engagements/{eid}/current-licenses", json={
        "sku_reference": "E3", "quantity_purchased": 110, "quantity_assigned": 100,
        "unit_price_paid_annual": 240, "unused_seats_answer": "NotNeeded"})
    html = client.get(f"/api/engagements/{eid}/readout.html").text
    assert "How the headline is timed" in html
    assert "Unused licences" in html and "not needed — counted as over-licensing" in html
    assert "(assumed)" in html                       # no Microsoft renewal date given
    wb = load_workbook(io.BytesIO(client.get(f"/api/engagements/{eid}/readout.xlsx").content))
    timing = list(wb["Headline timing"].iter_rows(values_only=True))
    assert timing[1][0] == "overlicensing" and timing[1][4] == "yes"
    assert list(wb["Unused licences"].iter_rows(values_only=True))[1][:2] == ("E3", 10)


def test_engagement_without_a_workshop_date_is_timed_from_its_creation(client):
    from app import models
    from app.db import SessionLocal

    eid, _ = _eng(client, "No Date Co")
    with SessionLocal() as db:
        eng = db.get(models.Engagement, eid)
        eng.workshop_date = None
        db.commit()
        created = eng.created_at.date()
    h = _headline(client, eid)
    assert h["workshop_date"] == created.isoformat()
